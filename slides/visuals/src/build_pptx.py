"""Assemble the ACD Perú 2026 deck on the official template.

Run: uv run --with python-pptx python slides/visuals/src/build_pptx.py <template.pptx>
Keeps the template's own title, ¿Preguntas?, event-close and ¡Gracias! slides (filled in), drops the
other example slides, and puts our rendered PNG posters in between as full-bleed slides.
"""
import copy
import pathlib
import sys

from pptx import Presentation
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

VIS = pathlib.Path(__file__).resolve().parent.parent
OUT = VIS.parent / "ARMed-and-Dangerous_ACD-Peru-2026.pptx"

# talk order
ORDER = ["02-pregunta", "03-graviton5", "04-chiste", "06-por-que-no", "05-nombre-instancia", "07-workloads",
         "07b-workloads-extra", "08-arquitectura", "09-como-medimos", "10-vcpu", "11-knee", "12-java", "13-go",
         "14-postgres", "15-mongo", "15a-inferencia", "15b-red", "16-perillas", "17-precio", "18-por-dolar",
         "19-rendimiento", "19a-arco", "19b-arco-por-dolar", "20-compatibilidad", "21-esfuerzo", "21b-fargate",
         "22-conclusiones", "23-recomendaciones", "24-aprendizajes", "25-casos-futuros", "26-siguientes-pasos"]
AFTER_QUESTIONS = ["28-feedback"]   # goes right after the template ¿Preguntas? slide
KEEP = {0: "title", 29: "questions", 30: "event-close", 31: "thanks"}   # template slide index -> role
NAMES = "Andrés Zeballos · Victor Herrera"
ROLES = "Solutions Architect · phData  |  Cloud Specialist · Caleidos"
PENDING = "POR-CONFIRMAR"
# slides that play their MP4 on entry (motion carries the idea); True = loop
VIDEO = {"05-nombre-instancia": False, "08-arquitectura": True, "10-vcpu": False, "11-knee": False,
         "20-compatibilidad": False, "21-esfuerzo": False}
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
AUTOPLAY = """<p:timing xmlns:p="%s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="%d" fill="hold"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" %sfill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def add_video(slide, name, w, h, loop):
    """Full-bleed MP4 that starts on slide entry; the PNG poster is the frame shown before/after."""
    mv = slide.shapes.add_movie(str(VIS / f"{name}.mp4"), 0, 0, w, h,
                                poster_frame_image=str(VIS / f"{name}.png"), mime_type="video/mp4")
    sld = slide._element
    for t in sld.findall(f"{{{P}}}timing"):
        sld.remove(t)
    dur = 9000 if loop else 6000
    sld.append(etree.fromstring(AUTOPLAY % (P, dur, mv.shape_id, 'repeatCount="indefinite" ' if loop else "", mv.shape_id)))


def number(slide, n, w, h):
    tb = slide.shapes.add_textbox(w - Emu(914400), h - Emu(420000), Emu(640080), Emu(300000))
    para = tb.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT
    r = para.add_run(); r.text = str(n)
    r.font.size = Pt(11); r.font.name = "Inter"; r.font.color.rgb = RGBColor(0x8C, 0x86, 0xB6)


def set_text(shape, text):
    """Replace a placeholder's text but keep the template's run formatting."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    run = first.runs[0] if first.runs else first.add_run()
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in first.runs[1:]:
        r._r.getparent().remove(r._r)
    run.text = text
    return run


def by_text(slide, prefix):
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith(prefix))


def fill_title(slide):
    t = by_text(slide, "Título de tu charla")
    run = set_text(t, "ARMed and Dangerous")
    p2 = copy.deepcopy(t.text_frame.paragraphs[0]._p)
    t.text_frame._txBody.append(p2)
    sub = t.text_frame.paragraphs[1].runs[0]
    sub.text = "lo que Graviton5 hace con tus workloads, medido en EKS"
    sub.font.size = Pt(18)   # 20 pt wrapped "EKS" alone onto a third line
    run.font.size = Pt(40)
    set_text(by_text(slide, "Nombre Apellido"), NAMES)
    set_text(by_text(slide, "Cargo"), ROLES)
    set_text(by_text(slide, "Track de tu charla"), "Cloud Architecture & Modernization").font.size = Pt(12)   # one line in its column
    set_text(by_text(slide, "Sala"), "Colaboratorio")
    set_text(by_text(slide, "00:00"), "15:05")


def fill_thanks(slide):
    set_text(by_text(slide, "Nombre Apellido"), NAMES)
    set_text(by_text(slide, "Cargo"), ROLES)
    t = by_text(slide, "in/usuario")
    set_text(t, "Andrés: in/andreszc · github.com/andrezc98")
    p2 = copy.deepcopy(t.text_frame.paragraphs[0]._p); t.text_frame._txBody.append(p2)
    t.text_frame.paragraphs[1].runs[0].text = "Victor: in/victor-herreraa · github.com/VotircH"
    for para in t.text_frame.paragraphs:
        para.runs[0].font.size = Pt(14)
    qr = next(sh for sh in slide.placeholders if not sh.has_text_frame)   # template "Tu QR" sample picture
    slide.shapes.add_picture(str(VIS / "assets/repo-qr.png"), qr.left, qr.top, qr.width, qr.height)
    qr._element.getparent().remove(qr._element)


def main(template):
    prs = Presentation(template)
    slides = list(prs.slides)
    fill_title(slides[0])
    fill_thanks(slides[31])
    # drop every template example slide we don't keep
    ids = prs.slides._sldIdLst
    for i, sld in reversed(list(enumerate(list(ids)))):
        if i not in KEEP:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
    blank = next(l for l in prs.slide_layouts if l.name == "DEFAULT")
    W, H = prs.slide_width, prs.slide_height
    for name in ORDER:
        s = prs.slides.add_slide(blank)
        if name in VIDEO:
            add_video(s, name, W, H, VIDEO[name])
        else:
            s.shapes.add_picture(str(VIS / f"{name}.png"), 0, 0, W, H)
    for name in AFTER_QUESTIONS:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(VIS / f"{name}.png"), 0, 0, W, H)
    fb = [ids[-1]]
    for x in fb:
        ids.remove(x)
    # move the closing template slides to the end: ¿Preguntas?, feedback, event close, ¡Gracias!
    q, close, thanks = ids[1], ids[2], ids[3]
    for x in (q, close, thanks):
        ids.remove(x)
    for x in (q, *fb, close, thanks):
        ids.append(x)
    # slide numbers bottom right, except the cover and the event-close slide
    total = len(prs.slides)
    for n, sl in enumerate(prs.slides, 1):
        if n not in (1, total - 1):
            number(sl, n, W, H)
    # unique part names (python-pptx reuses slideN.xml names of dropped slides)
    from pptx.opc.packuri import PackURI
    for i, sld in enumerate(prs.slides, 1):
        sld.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
    prs.save(OUT)
    print(OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main(sys.argv[1])
